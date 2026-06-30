"""
Intervento fotovoltaico: installazione impianto FV con calcolo producibilità via PVGIS.

Input richiesti (in parametri):
    lat (float)              : latitudine sito
    lon (float)              : longitudine sito
    superfici (list[dict])   : lista superfici FV
                               ciascuna: {id, slope, azimuth, n_pannelli, potenza_wp,
                                          pvtech='crystSi', mounting='free'}
    percorso_consumi (str)   : percorso file consumi quart'orari (CSV o Excel)
    anno_pvgis (int)         : anno simulazione PVGIS, default 2023
    loss (float)             : perdite sistema [%], default 14.0
    prezzo_kwh (float)       : prezzo acquisto energia [€/kWh], default 0.21
    pun_euro_kwh (float)     : prezzo cessione in rete (PUN) [€/kWh], default 0.12
    pod_selezionati (list)   : codici POD per cui dimensionare l'autoconsumo (None = tutti)
    fabbricati (list[str])   : fabbricati sotto il POD (descrittivo, per il report)
    edifici (list[str])      : edifici considerati per le superfici (descrittivo, per il report)

Modello economico (direttive di progetto):
    prezziario impianto = SEMPRE Marche (a prescindere dalla regione del sito)
    investimento = costo_impianto + 15% progettazione
    manutenzione annua = 1% di (costo_impianto + progettazione)
    flusso netto annuo = autoconsumo·prezzo + immissione·PUN − manutenzione
    vita utile = 20 anni, tasso di sconto = 6%
    nessun incentivo, nessun Certificato Bianco
"""

from __future__ import annotations
from typing import Any, Dict, List, Optional

from pydantic import BaseModel, Field

from core.energy_model import EnergyModel
from core.consumi_parser import parse_consumi_quart_orari, parse_consumi_excel_multi_pod
from interventions.base import InterventionResult
from interventions.economics import (
    VanScenario, calcola_van_scenario, _DISCOUNT_RATE,
)
from interventions.fotovoltaico_pvgis import get_producibilita_oraria
from interventions.fotovoltaico_prezzi import costo_impianto

NOME_INTERVENTO = "fotovoltaico"

_PREZZO_KWH_DEFAULT = 0.21
_PUN_DEFAULT = 0.12
_FATTORE_CO2 = 294.784 / 1000     # kg/kWh
_FATTORE_TEP = 0.000187           # tep/kWh
_QUOTA_PROGETTAZIONE = 0.15       # 15% del costo impianto
_QUOTA_MANUTENZIONE = 0.01        # 1% annuo di (impianto + progettazione)
_VITA_UTILE_ANNI = 20
# Il costo dell'impianto FV si valuta SEMPRE con il prezziario regionale delle Marche,
# indipendentemente dalla regione del sito (direttiva di progetto).
_REGIONE_PREZZIARIO_FV = "marche"


# ---------------------------------------------------------------------------
# Modelli dati input
# ---------------------------------------------------------------------------

class SuperficieFV(BaseModel):
    """Superficie fotovoltaica (falda, pensilina, ecc.)."""
    id: str = Field(description="Identificativo superficie (es. 'Tetto Sud')")
    slope: float = Field(description="Inclinazione da orizzontale [°]: 0=piano, 90=verticale")
    azimuth: float = Field(description="Orientamento PVGIS [°]: 0=sud, -90=est, +90=ovest")
    n_pannelli: int = Field(description="Numero di pannelli")
    potenza_pannello_wp: float = Field(description="Potenza singolo pannello [Wp]")
    pvtech: str = Field(default="crystSi", description="Tecnologia: crystSi|CIS|CdTe|Unknown")
    mounting: str = Field(default="free", description="Tipo montaggio: free|building")

    @property
    def peak_power_kwp(self) -> float:
        return self.n_pannelli * self.potenza_pannello_wp / 1000.0


# ---------------------------------------------------------------------------
# Modelli dati output
# ---------------------------------------------------------------------------

class FVRiga(BaseModel):
    """Producibilità per singola superficie."""
    superficie_id: str
    slope: float
    azimuth: float
    n_pannelli: int
    potenza_pannello_wp: float
    peak_power_kwp: float
    e_prodotta_kwh: float
    ore_equivalenti: float


class FVResult(InterventionResult):
    """Risultato completo dell'intervento fotovoltaico."""
    superfici_input: List[SuperficieFV]
    righe: List[FVRiga]

    # Dati impianto
    potenza_totale_kwp: float
    costo_per_kwp: float

    # Investimento (impianto + progettazione)
    costo_impianto_euro: float
    progettazione_euro: float
    investimento_totale_euro: float
    manutenzione_annua_euro: float
    vita_utile_anni: int

    # Producibilità
    e_prodotta_kwh: float
    ore_equivalenti_impianto: float

    # Autoconsumo
    e_autoconsumata_kwh: float
    e_immessa_kwh: float
    e_prelevata_kwh: float
    quota_autoconsumo: float

    # Consumo del/i POD selezionato/i
    consumo_pod_kwh: float
    pod_selezionati: List[str]

    # Economico annuo
    risparmio_acquisto_euro: float
    ricavo_immissione_euro: float
    risparmio_lordo_euro: float
    risparmio_totale_euro: float
    co2_evitata_kg: float
    tep_risparmiati: float

    # Benchmark VAN (il FV non usa incentivi né Certificati Bianchi)
    van_attualizzato: VanScenario
    van_semplice: VanScenario

    # Serie orarie (8760) per i grafici del report
    produzione_oraria: List[float] = Field(default_factory=list)
    consumo_orario: List[float] = Field(default_factory=list)

    # Metadati report (descrittivi)
    fabbricati: List[str] = Field(default_factory=list)
    edifici: List[str] = Field(default_factory=list)

    # Parametri usati
    prezzo_kwh: float
    pun_euro_kwh: float
    anno_pvgis: int


# ---------------------------------------------------------------------------
# Parsing tabella superfici
# ---------------------------------------------------------------------------

def _parse_superfici(superfici_raw: list[dict]) -> list[SuperficieFV]:
    result = []
    for d in superfici_raw:
        result.append(SuperficieFV(
            id=str(d.get("id", f"Superficie_{len(result)+1}")),
            slope=float(d["slope"]),
            azimuth=float(d["azimuth"]),
            n_pannelli=int(d["n_pannelli"]),
            potenza_pannello_wp=float(d["potenza_wp"]),
            pvtech=str(d.get("pvtech", "crystSi")),
            mounting=str(d.get("mounting", "free")),
        ))
    return result


def _parse_superfici_file(filepath: str) -> list[SuperficieFV]:
    from pathlib import Path
    import csv

    fp = Path(filepath)
    if not fp.exists():
        raise FileNotFoundError(f"File superfici non trovato: {filepath}")

    ext = fp.suffix.lower()
    rows: list[dict] = []

    if ext in (".xlsx", ".xls"):
        import openpyxl
        wb = openpyxl.load_workbook(fp, data_only=True, read_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            raise ValueError("File superfici vuoto")
        header = [str(c).strip().lower() if c else f"col{i}" for i, c in enumerate(all_rows[0])]
        for row in all_rows[1:]:
            if row and any(v is not None for v in row):
                rows.append(dict(zip(header, row)))
    else:
        with open(fp, newline="", encoding="utf-8-sig") as f:
            sample = f.read(2048)
            f.seek(0)
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            reader = csv.DictReader(f, dialect=dialect)
            rows = [r for r in reader]

    superfici = []
    for i, row in enumerate(rows):
        norm = {k.strip().lower(): v for k, v in row.items() if v not in (None, "")}

        def get(*keys):
            for k in keys:
                if k in norm:
                    return norm[k]
            return None

        id_val = get("id", "nome", "superficie") or f"Superficie_{i+1}"
        slope = get("slope", "tilt", "inclinazione")
        azimuth = get("azimuth", "aspect", "orientamento")
        n = get("n_pannelli", "n pannelli", "pannelli", "n")
        pwr = get("potenza_wp", "potenza wp", "wp", "potenza_pannello_wp", "potenza")
        pvtech = get("pvtech", "tecnologia") or "crystSi"
        mounting = get("mounting", "montaggio") or "free"

        if any(v is None for v in [slope, azimuth, n, pwr]):
            raise ValueError(
                f"Riga {i+2} del file superfici: campi obbligatori mancanti "
                f"(slope, azimuth, n_pannelli, potenza_wp). Trovato: {dict(norm)}"
            )

        superfici.append(SuperficieFV(
            id=str(id_val),
            slope=float(slope),
            azimuth=float(azimuth),
            n_pannelli=int(float(n)),
            potenza_pannello_wp=float(pwr),
            pvtech=str(pvtech),
            mounting=str(mounting),
        ))

    return superfici


# ---------------------------------------------------------------------------
# Calcolo autoconsumo
# ---------------------------------------------------------------------------

def _calcola_autoconsumo(
    produzione: list[float],
    consumo: list[float],
) -> tuple[float, float, float]:
    auto = imm = prel = 0.0
    for p, c in zip(produzione, consumo):
        auto += min(p, c)
        imm  += max(0.0, p - c)
        prel += max(0.0, c - p)
    return auto, imm, prel


# ---------------------------------------------------------------------------
# Funzione principale
# ---------------------------------------------------------------------------

def calcola(model: EnergyModel, parametri: Dict[str, Any]) -> FVResult:
    """
    Calcola il benchmark dell'intervento fotovoltaico.

    Se parametri contiene '_produzione_oraria_precalcolata' (lista 8760 valori)
    e '_pvgis_righe' (lista dict per superficie), le chiamate PVGIS vengono saltate
    e si usano i dati gia' calcolati (es. dalla pipeline GitHub Actions).
    """
    lat = float(parametri["lat"])
    lon = float(parametri["lon"])
    anno = int(parametri.get("anno_pvgis", 2023))
    loss = float(parametri.get("loss", 14.0))
    prezzo_kwh = float(parametri.get("prezzo_kwh", _PREZZO_KWH_DEFAULT))
    pun = float(parametri.get("pun_euro_kwh", _PUN_DEFAULT))

    pod_sel_raw = parametri.get("pod_selezionati") or []
    pod_selezionati = [str(p).strip() for p in pod_sel_raw if str(p).strip()]

    fabbricati = [str(x) for x in (parametri.get("fabbricati") or [])]
    edifici = [str(x) for x in (parametri.get("edifici") or [])]

    if "superfici" in parametri:
        superfici = _parse_superfici(parametri["superfici"])
    elif "percorso_superfici" in parametri:
        superfici = _parse_superfici_file(str(parametri["percorso_superfici"]))
    else:
        raise ValueError("parametri deve contenere 'superfici' (list) o 'percorso_superfici' (path)")

    if not superfici:
        raise ValueError("Nessuna superficie FV definita")

    if "percorso_consumi" not in parametri:
        raise ValueError("parametri deve contenere 'percorso_consumi'")
    percorso_consumi = str(parametri["percorso_consumi"])
    if percorso_consumi.lower().endswith((".xlsx", ".xls")):
        try:
            import openpyxl
            wb_check = openpyxl.load_workbook(percorso_consumi, data_only=True, read_only=True)
            foglio_multi = "Input POD orario ATTIVA"
            if foglio_multi in wb_check.sheetnames:
                col_prima_pod = int(parametri.get("col_prima_pod", 130))
                consumo_orario = parse_consumi_excel_multi_pod(
                    percorso_consumi, nome_foglio=foglio_multi, col_prima_pod=col_prima_pod,
                    pod_selezionati=pod_selezionati or None,
                )
            else:
                consumo_orario = parse_consumi_quart_orari(percorso_consumi)
        except Exception:
            consumo_orario = parse_consumi_quart_orari(percorso_consumi)
    else:
        consumo_orario = parse_consumi_quart_orari(percorso_consumi)

    # PVGIS: usa dati precalcolati se disponibili, altrimenti chiama l'API
    produzione_totale: list[float] = [0.0] * 8760
    righe: list[FVRiga] = []

    if "_produzione_oraria_precalcolata" in parametri:
        produzione_totale = list(parametri["_produzione_oraria_precalcolata"])[:8760]
        for raw_riga in parametri.get("_pvgis_righe", []):
            sup_match = next((s for s in superfici if s.id == raw_riga["superficie_id"]), None)
            if sup_match is None:
                continue
            righe.append(FVRiga(
                superficie_id=raw_riga["superficie_id"],
                slope=raw_riga["slope"],
                azimuth=raw_riga["azimuth"],
                n_pannelli=raw_riga["n_pannelli"],
                potenza_pannello_wp=raw_riga["potenza_pannello_wp"],
                peak_power_kwp=round(raw_riga["peak_power_kwp"], 3),
                e_prodotta_kwh=round(raw_riga["e_prodotta_kwh"], 1),
                ore_equivalenti=round(raw_riga["ore_equivalenti"], 1),
            ))
    else:
        import time
        for i, sup in enumerate(superfici):
            if i > 0:
                time.sleep(0.5)
            prod_sup = get_producibilita_oraria(
                lat=lat, lon=lon,
                peak_power_kwp=sup.peak_power_kwp,
                slope=sup.slope, aspect=sup.azimuth,
                anno=anno, loss=loss,
                pvtech=sup.pvtech, mounting=sup.mounting,
            )
            n = min(len(prod_sup), 8760)
            for h in range(n):
                produzione_totale[h] += prod_sup[h]
            e_sup = sum(prod_sup[:n])
            h_eq = e_sup / sup.peak_power_kwp if sup.peak_power_kwp > 0 else 0.0
            righe.append(FVRiga(
                superficie_id=sup.id,
                slope=sup.slope,
                azimuth=sup.azimuth,
                n_pannelli=sup.n_pannelli,
                potenza_pannello_wp=sup.potenza_pannello_wp,
                peak_power_kwp=round(sup.peak_power_kwp, 3),
                e_prodotta_kwh=round(e_sup, 1),
                ore_equivalenti=round(h_eq, 1),
            ))

    potenza_totale = sum(s.peak_power_kwp for s in superfici)
    e_prodotta = sum(produzione_totale)
    ore_eq_impianto = e_prodotta / potenza_totale if potenza_totale > 0 else 0.0

    e_auto, e_imm, e_prel = _calcola_autoconsumo(produzione_totale, consumo_orario)
    quota_auto = e_auto / e_prodotta if e_prodotta > 0 else 0.0
    consumo_pod = sum(consumo_orario)

    costo_imp = costo_impianto(potenza_totale, _REGIONE_PREZZIARIO_FV)
    progettazione = costo_imp * _QUOTA_PROGETTAZIONE
    investimento_totale = costo_imp + progettazione
    costo_unitario = costo_imp / potenza_totale if potenza_totale > 0 else 0.0

    manutenzione_annua = investimento_totale * _QUOTA_MANUTENZIONE

    risp_acquisto = e_auto * prezzo_kwh
    ricavo_imm = e_imm * pun
    risp_lordo = risp_acquisto + ricavo_imm
    risp_netto = risp_lordo - manutenzione_annua
    co2_evitata = e_auto * _FATTORE_CO2
    tep_risparmiati = e_auto * _FATTORE_TEP

    scenari_config = [
        ("VAN Attualizzato", _DISCOUNT_RATE),
        ("VAN Semplice",     0.0),
    ]
    van_results: list[VanScenario] = []
    for nome, dr in scenari_config:
        van_results.append(calcola_van_scenario(
            investimento=investimento_totale,
            risparmio_annuo_euro=risp_netto,
            risparmio_annuo_ee_kwh=e_auto,
            incentivo_annuo=0.0,
            discount_rate=dr,
            include_incentivi=False,
            nome=nome,
            vita_anni=_VITA_UTILE_ANNI,
        ))

    return FVResult(
        nome_intervento=NOME_INTERVENTO,
        risparmio_annuo_kwh=round(e_auto, 1),
        costo_stimato_euro=round(investimento_totale, 2),
        payback_anni=round(van_results[0].tr, 2),
        dati_grafico={
            "e_prodotta_kwh": e_prodotta,
            "e_autoconsumata_kwh": e_auto,
            "e_immessa_kwh": e_imm,
            "quota_autoconsumo": quota_auto,
            "risparmio_totale_euro": risp_netto,
            "co2_evitata_kg": co2_evitata,
            "investimento": investimento_totale,
            "van_attualizzato": van_results[0].van,
            "van_semplice": van_results[1].van,
        },
        note=[],
        superfici_input=superfici,
        righe=righe,
        potenza_totale_kwp=round(potenza_totale, 3),
        costo_per_kwp=round(costo_unitario, 2),
        costo_impianto_euro=round(costo_imp, 2),
        progettazione_euro=round(progettazione, 2),
        investimento_totale_euro=round(investimento_totale, 2),
        manutenzione_annua_euro=round(manutenzione_annua, 2),
        vita_utile_anni=_VITA_UTILE_ANNI,
        e_prodotta_kwh=round(e_prodotta, 1),
        ore_equivalenti_impianto=round(ore_eq_impianto, 1),
        e_autoconsumata_kwh=round(e_auto, 1),
        e_immessa_kwh=round(e_imm, 1),
        e_prelevata_kwh=round(e_prel, 1),
        quota_autoconsumo=round(quota_auto, 4),
        consumo_pod_kwh=round(consumo_pod, 1),
        pod_selezionati=pod_selezionati,
        risparmio_acquisto_euro=round(risp_acquisto, 2),
        ricavo_immissione_euro=round(ricavo_imm, 2),
        risparmio_lordo_euro=round(risp_lordo, 2),
        risparmio_totale_euro=round(risp_netto, 2),
        co2_evitata_kg=round(co2_evitata, 1),
        tep_risparmiati=round(tep_risparmiati, 4),
        van_attualizzato=van_results[0],
        van_semplice=van_results[1],
        produzione_oraria=[round(x, 4) for x in produzione_totale],
        consumo_orario=[round(x, 4) for x in consumo_orario],
        fabbricati=fabbricati,
        edifici=edifici,
        prezzo_kwh=prezzo_kwh,
        pun_euro_kwh=pun,
        anno_pvgis=anno,
    )
