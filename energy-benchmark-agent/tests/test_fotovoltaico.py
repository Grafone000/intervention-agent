"""
Test del modulo fotovoltaico.
PVGIS API viene mockato — nessuna chiamata di rete.
"""

from __future__ import annotations
import csv
import tempfile
from pathlib import Path
from unittest.mock import patch

import pytest

from interventions.fotovoltaico import (
    SuperficieFV, calcola, _calcola_autoconsumo, _parse_superfici,
)
from interventions.fotovoltaico_prezzi import costo_per_kwp, costo_impianto
from interventions.economics import calcola_cb, calcola_van_scenario


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _produzione_sintetica(kwh_per_ora_diurna: float = 5.0) -> list[float]:
    """8760 valori: 0 di notte (0-5 e 20-23), kwh_per_ora_diurna di giorno."""
    out = []
    for h in range(8760):
        ora = h % 24
        out.append(kwh_per_ora_diurna if 6 <= ora < 20 else 0.0)
    return out


def _consumi_sintetici(kwh_per_ora: float = 3.0) -> list[float]:
    return [kwh_per_ora] * 8760


def _scrivi_consumi_quart_orari(path: Path, kwh_per_quarto: float = 0.75) -> None:
    """Scrive 35040 righe CSV (4 × 8760 = un anno quart'orario)."""
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["timestamp", "consumo_kwh"])
        for i in range(35040):
            w.writerow([f"2023-01-01 {i:05d}", kwh_per_quarto])


# ---------------------------------------------------------------------------
# SuperficieFV
# ---------------------------------------------------------------------------

def test_peak_power_kwp():
    sup = SuperficieFV(id="A", slope=30, azimuth=0, n_pannelli=20, potenza_pannello_wp=400)
    assert abs(sup.peak_power_kwp - 8.0) < 1e-6


def test_parse_superfici_da_dict():
    raw = [{"id": "Tetto", "slope": 30, "azimuth": 0, "n_pannelli": 10, "potenza_wp": 400}]
    sups = _parse_superfici(raw)
    assert len(sups) == 1
    assert sups[0].peak_power_kwp == 4.0


# ---------------------------------------------------------------------------
# Autoconsumo
# ---------------------------------------------------------------------------

def test_autoconsumo_produzione_nulla():
    prod = [0.0] * 8760
    cons = [3.0] * 8760
    auto, imm, prel = _calcola_autoconsumo(prod, cons)
    assert auto == 0.0
    assert imm == 0.0
    assert abs(prel - 3.0 * 8760) < 0.01


def test_autoconsumo_produzione_eccedente():
    prod = [10.0] * 8760
    cons = [3.0] * 8760
    auto, imm, prel = _calcola_autoconsumo(prod, cons)
    assert abs(auto - 3.0 * 8760) < 0.01
    assert abs(imm - 7.0 * 8760) < 0.01
    assert prel == 0.0


def test_autoconsumo_bilanciato():
    prod = [3.0] * 8760
    cons = [3.0] * 8760
    auto, imm, prel = _calcola_autoconsumo(prod, cons)
    assert abs(auto - 3.0 * 8760) < 0.01
    assert imm == 0.0
    assert prel == 0.0


def test_autoconsumo_diurno():
    """Di giorno produzione > consumo, di notte produzione = 0."""
    prod = _produzione_sintetica(5.0)   # 5 kWh ore 6-19
    cons = _consumi_sintetici(3.0)       # 3 kWh sempre
    auto, imm, prel = _calcola_autoconsumo(prod, cons)
    # Di giorno (14h): auto=3, imm=2; di notte (10h): auto=0, prel=3
    ore_diurne = 14 * 365
    ore_notturne = 10 * 365
    assert abs(auto - 3.0 * ore_diurne) < 1.0
    assert abs(imm - 2.0 * ore_diurne) < 1.0
    assert abs(prel - 3.0 * ore_notturne) < 1.0


# ---------------------------------------------------------------------------
# Prezziario
# ---------------------------------------------------------------------------

def test_prezzo_piccolo_impianto():
    assert costo_per_kwp(5.0) == 1500.0

def test_prezzo_medio_impianto():
    assert costo_per_kwp(15.0) == 1250.0

def test_prezzo_grande_impianto():
    assert costo_per_kwp(200.0) == 800.0

def test_prezzo_sardegna():
    assert costo_per_kwp(5.0, "Sardegna") == 2284.46

def test_prezzo_regione_sconosciuta_usa_mercato():
    assert costo_per_kwp(5.0, "Sicilia") == 1500.0

def test_costo_impianto_totale():
    c = costo_impianto(10.0)
    assert abs(c - 10.0 * 1250.0) < 0.01


# ---------------------------------------------------------------------------
# Economics
# ---------------------------------------------------------------------------

def test_cb_tep_arrotondati():
    cb = calcola_cb(14.7)
    assert cb.tep_arrotondati == 15
    assert abs(cb.incentivo_annuo - 15 * 250.0) < 0.01

def test_cb_zero():
    cb = calcola_cb(0.3)
    assert cb.tep_arrotondati == 0
    assert cb.incentivo_annuo == 0.0

def test_van_con_incentivi_maggiore_senza():
    van_ci = calcola_van_scenario(
        investimento=50000, risparmio_annuo_euro=8000,
        risparmio_annuo_ee_kwh=40000, incentivo_annuo=3000,
        discount_rate=0.06, include_incentivi=True, nome="CI"
    )
    van_si = calcola_van_scenario(
        investimento=50000, risparmio_annuo_euro=8000,
        risparmio_annuo_ee_kwh=40000, incentivo_annuo=3000,
        discount_rate=0.06, include_incentivi=False, nome="SI"
    )
    assert van_ci.van >= van_si.van

def test_van_fc_anno8_uguale_van():
    van = calcola_van_scenario(
        investimento=10000, risparmio_annuo_euro=2000,
        risparmio_annuo_ee_kwh=10000, incentivo_annuo=0,
        discount_rate=0.06, include_incentivi=False, nome="test"
    )
    assert abs(van.van - van.flussi_cassa[-1]) < 0.1


# ---------------------------------------------------------------------------
# calcola() — integrazione con PVGIS mockato
# ---------------------------------------------------------------------------

@pytest.fixture
def consumi_file(tmp_path):
    """File consumi quart'orari sintetici: 3 kWh/ora → 0.75 kWh per quarto."""
    p = tmp_path / "consumi.csv"
    _scrivi_consumi_quart_orari(p, kwh_per_quarto=0.75)
    return str(p)


@pytest.fixture
def mock_pvgis():
    """Mocka PVGIS: restituisce 5 kWh/h ore 6-19, 0 altrove."""
    prod = _produzione_sintetica(5.0)
    with patch("interventions.fotovoltaico.get_producibilita_oraria", return_value=prod):
        yield


def test_calcola_tipo_risultato(mock_pvgis, consumi_file):
    from core.energy_model import EnergyModel
    from interventions.fotovoltaico import FVResult
    model = EnergyModel(nome_progetto="test", utenze=[])
    params = {
        "lat": 40.65, "lon": 8.91,
        "superfici": [{"id": "T", "slope": 30, "azimuth": 0, "n_pannelli": 20, "potenza_wp": 400}],
        "percorso_consumi": consumi_file,
    }
    result = calcola(model, params)
    assert isinstance(result, FVResult)
    assert result.nome_intervento == "fotovoltaico"


def test_calcola_produzione_positiva(mock_pvgis, consumi_file):
    from core.energy_model import EnergyModel
    model = EnergyModel(nome_progetto="test", utenze=[])
    params = {
        "lat": 40.65, "lon": 8.91,
        "superfici": [{"id": "T", "slope": 30, "azimuth": 0, "n_pannelli": 20, "potenza_wp": 400}],
        "percorso_consumi": consumi_file,
    }
    result = calcola(model, params)
    assert result.e_prodotta_kwh > 0
    assert result.e_autoconsumata_kwh > 0
    assert result.quota_autoconsumo > 0


def test_calcola_quota_autoconsumo_range(mock_pvgis, consumi_file):
    from core.energy_model import EnergyModel
    model = EnergyModel(nome_progetto="test", utenze=[])
    params = {
        "lat": 40.65, "lon": 8.91,
        "superfici": [{"id": "T", "slope": 30, "azimuth": 0, "n_pannelli": 20, "potenza_wp": 400}],
        "percorso_consumi": consumi_file,
    }
    result = calcola(model, params)
    assert 0.0 <= result.quota_autoconsumo <= 1.0


def test_calcola_van_attualizzato_minore_semplice(mock_pvgis, consumi_file):
    """Senza incentivi, il VAN attualizzato (sconto 6%) è ≤ del VAN semplice (sconto 0%)."""
    from core.energy_model import EnergyModel
    model = EnergyModel(nome_progetto="test", utenze=[])
    params = {
        "lat": 40.65, "lon": 8.91,
        "superfici": [{"id": "T", "slope": 30, "azimuth": 0, "n_pannelli": 20, "potenza_wp": 400}],
        "percorso_consumi": consumi_file,
    }
    result = calcola(model, params)
    assert result.van_attualizzato.van <= result.van_semplice.van


def test_calcola_tir_solo_attualizzato(mock_pvgis, consumi_file):
    from core.energy_model import EnergyModel
    model = EnergyModel(nome_progetto="test", utenze=[])
    params = {
        "lat": 40.65, "lon": 8.91,
        "superfici": [{"id": "T", "slope": 30, "azimuth": 0, "n_pannelli": 20, "potenza_wp": 400}],
        "percorso_consumi": consumi_file,
    }
    result = calcola(model, params)
    assert result.van_semplice.tir is None
    assert result.van_attualizzato.tir is not None


def test_calcola_progettazione_15pct(mock_pvgis, consumi_file):
    from core.energy_model import EnergyModel
    model = EnergyModel(nome_progetto="test", utenze=[])
    params = {
        "lat": 40.65, "lon": 8.91,
        "superfici": [{"id": "T", "slope": 30, "azimuth": 0, "n_pannelli": 20, "potenza_wp": 400}],
        "percorso_consumi": consumi_file,
    }
    result = calcola(model, params)
    # progettazione = 15% del costo impianto
    assert abs(result.progettazione_euro - result.costo_impianto_euro * 0.15) < 0.05
    # investimento totale = impianto + progettazione
    assert abs(result.investimento_totale_euro
               - (result.costo_impianto_euro + result.progettazione_euro)) < 0.05


def test_calcola_manutenzione_1pct(mock_pvgis, consumi_file):
    from core.energy_model import EnergyModel
    model = EnergyModel(nome_progetto="test", utenze=[])
    params = {
        "lat": 40.65, "lon": 8.91,
        "superfici": [{"id": "T", "slope": 30, "azimuth": 0, "n_pannelli": 20, "potenza_wp": 400}],
        "percorso_consumi": consumi_file,
    }
    result = calcola(model, params)
    # manutenzione = 1% di (impianto + progettazione)
    assert abs(result.manutenzione_annua_euro - result.investimento_totale_euro * 0.01) < 0.05
    # risparmio netto = lordo - manutenzione
    assert abs(result.risparmio_totale_euro
               - (result.risparmio_lordo_euro - result.manutenzione_annua_euro)) < 0.05


def test_calcola_vita_utile_20_anni(mock_pvgis, consumi_file):
    from core.energy_model import EnergyModel
    model = EnergyModel(nome_progetto="test", utenze=[])
    params = {
        "lat": 40.65, "lon": 8.91,
        "superfici": [{"id": "T", "slope": 30, "azimuth": 0, "n_pannelli": 20, "potenza_wp": 400}],
        "percorso_consumi": consumi_file,
    }
    result = calcola(model, params)
    assert result.vita_utile_anni == 20
    assert len(result.van_attualizzato.flussi_cassa) == 20
    assert len(result.van_semplice.flussi_cassa) == 20


def test_calcola_serie_orarie_salvate(mock_pvgis, consumi_file):
    from core.energy_model import EnergyModel
    model = EnergyModel(nome_progetto="test", utenze=[])
    params = {
        "lat": 40.65, "lon": 8.91,
        "superfici": [{"id": "T", "slope": 30, "azimuth": 0, "n_pannelli": 20, "potenza_wp": 400}],
        "percorso_consumi": consumi_file,
    }
    result = calcola(model, params)
    assert len(result.produzione_oraria) == 8760
    assert len(result.consumo_orario) == 8760


def _scrivi_excel_multi_pod(path: Path, pods: list[str], valori: dict[str, float],
                            col_prima_pod: int = 1, n_ore: int = 8760) -> None:
    """Crea un Excel con foglio 'Input POD orario ATTIVA': header POD + n_ore righe costanti."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Input POD orario ATTIVA"
    for j, pod in enumerate(pods):
        ws.cell(row=1, column=col_prima_pod + j, value=pod)
    for r in range(n_ore):
        for j, pod in enumerate(pods):
            ws.cell(row=2 + r, column=col_prima_pod + j, value=valori[pod])
    wb.save(path)


def test_lista_pod_disponibili(tmp_path):
    from core.consumi_parser import lista_pod_disponibili
    p = tmp_path / "multi.xlsx"
    _scrivi_excel_multi_pod(p, ["POD_A", "POD_B", "POD_C"],
                            {"POD_A": 1.0, "POD_B": 2.0, "POD_C": 3.0}, n_ore=10)
    assert lista_pod_disponibili(p, col_prima_pod=1) == ["POD_A", "POD_B", "POD_C"]


def test_selezione_pod_somma_solo_richiesti(tmp_path):
    from core.consumi_parser import parse_consumi_excel_multi_pod
    p = tmp_path / "multi.xlsx"
    _scrivi_excel_multi_pod(p, ["POD_A", "POD_B", "POD_C"],
                            {"POD_A": 1.0, "POD_B": 2.0, "POD_C": 3.0})
    # solo POD_A + POD_C → 4.0/ora
    orari = parse_consumi_excel_multi_pod(p, col_prima_pod=1, pod_selezionati=["POD_A", "POD_C"])
    assert len(orari) == 8760
    assert abs(orari[0] - 4.0) < 1e-9


def test_selezione_pod_tutti_se_none(tmp_path):
    from core.consumi_parser import parse_consumi_excel_multi_pod
    p = tmp_path / "multi.xlsx"
    _scrivi_excel_multi_pod(p, ["POD_A", "POD_B", "POD_C"],
                            {"POD_A": 1.0, "POD_B": 2.0, "POD_C": 3.0})
    orari = parse_consumi_excel_multi_pod(p, col_prima_pod=1, pod_selezionati=None)
    assert abs(orari[0] - 6.0) < 1e-9


def test_selezione_pod_inesistente_solleva(tmp_path):
    from core.consumi_parser import parse_consumi_excel_multi_pod
    p = tmp_path / "multi.xlsx"
    _scrivi_excel_multi_pod(p, ["POD_A"], {"POD_A": 1.0})
    with pytest.raises(ValueError):
        parse_consumi_excel_multi_pod(p, col_prima_pod=1, pod_selezionati=["POD_X"])


def test_report_fotovoltaico_docx(mock_pvgis, consumi_file, tmp_path):
    """La generazione del report .docx con sezione FV non deve sollevare eccezioni."""
    from core.energy_model import EnergyModel
    from reporting.build_report import costruisci_relazione
    model = EnergyModel(nome_progetto="test", utenze=[])
    params = {
        "lat": 40.65, "lon": 8.91,
        "superfici": [{"id": "FA18", "slope": 30, "azimuth": 44, "n_pannelli": 50, "potenza_wp": 450}],
        "percorso_consumi": consumi_file,
        "pod_selezionati": ["IT001E000001"],
        "fabbricati": ["Edificio A", "Edificio B"],
        "edifici": ["Edificio A"],
    }
    result = calcola(model, params)
    out = tmp_path / "report_fv.docx"
    path = costruisci_relazione([result], str(out))
    assert Path(path).exists()
    assert Path(path).stat().st_size > 0
